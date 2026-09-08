"""Выгрузка договора в PDF и Word (.docx)."""
import io

from django.template.loader import render_to_string

from apps.documents.pdf import render_pdf

from .contracts import build_contract_blocks


def contract_filename(contract, ext):
    number = (contract.number or "б-н").replace("/", "-")
    return f"Договор_{number}.{ext}"


def contract_pdf(contract, request=None):
    """PDF договора (тем же движком, что счета и накладные)."""
    html = render_to_string(
        "partners/print/contract.html",
        {"contract": contract, "blocks": build_contract_blocks(contract)},
        request=request,
    )
    base_url = request.build_absolute_uri("/") if request else None
    pdf, _engine = render_pdf(html, base_url=base_url)
    return pdf


def contract_docx(contract):
    """Word-версия договора: заголовки по центру, пункты с выключкой
    по ширине, реквизиты сторон — таблицей в две колонки."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    blocks = build_contract_blocks(contract)
    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(3)

    def paragraph(text="", bold=False, align=None, indent=None, hanging=None):
        para = document.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        if align is not None:
            para.alignment = align
        fmt = para.paragraph_format
        if indent is not None:
            fmt.first_line_indent = indent
        if hanging is not None:
            fmt.left_indent = hanging
            fmt.first_line_indent = -hanging
        return para

    if not blocks:
        paragraph("Договор не заполнен")
    else:
        paragraph(f"{blocks['title']} № {blocks['number']}", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        paragraph()

        # Место слева, дата справа — табуляцией по правому краю страницы
        meta = document.add_paragraph()
        meta.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_ALIGN_PARAGRAPH.RIGHT)
        meta.add_run(f"г. {blocks['place']}\t{blocks['date']}")
        paragraph()

        for text in blocks["intro"]:
            paragraph(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=Cm(1.25))

        for section in blocks["sections"]:
            paragraph()
            paragraph(f"{section['number']}. {section['title'].upper()}", bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
            for item in section["items"]:
                paragraph(f"{item['number']}. {item['text']}",
                          align=WD_ALIGN_PARAGRAPH.JUSTIFY, hanging=Cm(1.25))

        paragraph()
        paragraph(f"{blocks['requisites_number']}. АДРЕСА, РЕКВИЗИТЫ И ПОДПИСИ СТОРОН",
                  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        if blocks["parties"]:
            # Реквизиты и подписи — разными строками таблицы, иначе подписи
            # съезжают из-за разной длины реквизитов сторон.
            table = document.add_table(rows=2, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, party in zip(table.rows[0].cells, blocks["parties"]):
                cell.paragraphs[0].add_run(party["role"]).bold = True
                cell.add_paragraph().add_run(party["name"]).bold = True
                for label, value in party["rows"]:
                    cell.add_paragraph(f"{label}: {value}")
            for cell, party in zip(table.rows[1].cells, blocks["parties"]):
                cell.paragraphs[0].text = ""
                cell.add_paragraph("_________________ / "
                                   f"{party['signer'] or '________________'} /")
                cell.add_paragraph("М.П.")
        else:
            for line in (blocks["outro_text"] or "").split("\n"):
                paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
